"""QR-Codes für die Kartensätze erzeugen und (optional) ans Ende der Ich-kann-Liste setzen.

Aufruf (im Hauptverzeichnis des Repositorys):
    pip install segno python-docx zxing-cpp pillow
    python werkzeuge/qr_erzeugen.py kb1-diabetes                          # qr/qr_kb1-diabetes.png
    python werkzeuge/qr_erzeugen.py kb1-diabetes --klasse 3bfp2-1         # qr/qr_kb1-diabetes_3bfp2-1.png
    python werkzeuge/qr_erzeugen.py kb1-diabetes --klasse 3bfp2-1 Liste.docx   # + Liste_mit_QR_3BFP2-1.docx
    python werkzeuge/qr_erzeugen.py --sammelblatt QR-Codes_Klassen.docx   # alle Kartensätze × alle Klassen

- Liest Besitzer und Namen des Repositorys aus dem Git-Remote „origin“;
  die Adresse ist https://<besitzer>.github.io/<repository>/<ordner>/ (mit Klasse: …/?klasse=<kennung>).
- Klassen und Titel der Kartensätze stehen in docs/app/einstellungen.js.
- Erzeugt die QR-Codes (Fehlerkorrektur M, Ruhezone 4 Module, mindestens 600 × 600 px)
  und prüft, dass sie sich dekodieren lassen.
- Mit einer .docx-Datei: hängt unten einen Kasten mit QR-Code (3,2 cm) und kurzem Text an
  und speichert das Ergebnis als „…_mit_QR….docx“. Das Original bleibt unverändert.
- Sammelblatt: ein Word-Dokument mit allen Kästen zum Herauskopieren in die Ich-kann-Listen.
"""
import argparse
import re
import subprocess
import sys
from pathlib import Path

import segno
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

WURZEL = Path(__file__).resolve().parent.parent


def repo_aus_remote():
    url = subprocess.check_output(["git", "-C", str(WURZEL), "remote", "get-url", "origin"], text=True).strip()
    m = re.search(r"github\.com[/:]([^/]+)/([^/]+?)(?:\.git)?/?$", url)
    if not m:
        sys.exit(f"Git-Remote nicht erkannt: {url}")
    return m.group(1), m.group(2)


def qr_erzeugen(adresse, datei):
    qr = segno.make(adresse, error="m", micro=False)
    breite_module = qr.symbol_size(scale=1, border=4)[0]
    massstab = -(-600 // breite_module)  # aufrunden, damit mindestens 600 px
    qr.save(str(datei), kind="png", scale=massstab, border=4, dark="#000000", light="#FFFFFF")


def qr_pruefen(datei, erwartet):
    try:
        import zxingcpp
        from PIL import Image
    except ImportError:
        print("  (zxing-cpp/Pillow nicht installiert – Dekodierprüfung übersprungen)")
        return
    bild = Image.open(datei)
    ergebnisse = zxingcpp.read_barcodes(bild)
    texte = [e.text for e in ergebnisse]
    if texte != [erwartet]:
        sys.exit(f"QR-Prüfung fehlgeschlagen für {datei}: {texte}")
    print(f"  dekodiert: {texte[0]}  ({bild.size[0]} × {bild.size[1]} px, Fehlerkorrektur {ergebnisse[0].ec_level})")


def zelle_rahmen(zelle, farbe="3F6C72"):
    tc_pr = zelle._tc.get_or_add_tcPr()
    raender = OxmlElement("w:tcBorders")
    for seite in ("top", "left", "bottom", "right"):
        r = OxmlElement(f"w:{seite}")
        r.set(qn("w:val"), "single")
        r.set(qn("w:sz"), "12")
        r.set(qn("w:color"), farbe)
        raender.append(r)
    tc_pr.append(raender)


def einstellungen_lesen():
    """Klassen und Titel der Kartensätze aus docs/app/einstellungen.js."""
    text = (WURZEL / "docs" / "app" / "einstellungen.js").read_text(encoding="utf-8")
    klassen = dict(re.findall(r"\{\s*kennung:\s*'([a-z0-9-]+)',\s*name:\s*'([^']+)'\s*\}", text))
    titel = dict(re.findall(r"\{\s*kennung:\s*'([a-z0-9-]+)',\s*ordner:\s*'[^']+',\s*titel:\s*'([^']+)'", text))
    return klassen, titel


def kasten_anhaengen(doc, qr_datei, adresse, klassenname=None):
    tabelle = doc.add_table(rows=1, cols=2)
    tabelle.alignment = WD_TABLE_ALIGNMENT.CENTER
    tabelle.autofit = False
    links, rechts = tabelle.rows[0].cells
    links.width, rechts.width = Cm(3.8), Cm(11.5)
    for z in (links, rechts):
        zelle_rahmen(z)
        z.vertical_alignment = 1  # Mitte

    p = links.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(qr_datei), width=Cm(3.2))

    p = rechts.paragraphs[0]
    r = p.add_run("Lernkarten fürs Smartphone" + (f" – Klasse {klassenname}" if klassenname else ""))
    r.bold = True
    r.font.size = Pt(13)
    r.font.color.rgb = RGBColor(0x25, 0x43, 0x4B)
    rechts.add_paragraph(
        "QR-Code scannen und alle Kompetenzen dieser Liste als Karteikarten lernen: "
        "vorne die Kompetenz, hinten die Beispielantwort. Ohne Anmeldung – "
        "Ihr Lernstand bleibt auf Ihrem Handy."
    )
    r = rechts.add_paragraph().add_run(adresse.replace("https://", "").rstrip("/"))
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor(0x3F, 0x6C, 0x72)


def in_liste_einsetzen(docx_pfad, qr_datei, adresse, klassenname=None):
    doc = Document(docx_pfad)
    doc.add_paragraph()
    kasten_anhaengen(doc, qr_datei, adresse, klassenname)
    doc.add_paragraph()
    zusatz = "_" + klassenname.replace("/", "-") if klassenname else ""
    ziel = Path(docx_pfad).with_name(Path(docx_pfad).stem + "_mit_QR" + zusatz + ".docx")
    doc.save(ziel)
    return ziel


def qr_fuer(ordner, klasse=None):
    besitzer, repo = repo_aus_remote()
    adresse = f"https://{besitzer.lower()}.github.io/{repo}/{ordner}/" + (f"?klasse={klasse}" if klasse else "")
    qr_datei = WURZEL / "qr" / (f"qr_{ordner}_{klasse}.png" if klasse else f"qr_{ordner}.png")
    qr_datei.parent.mkdir(exist_ok=True)
    print(f"Adresse: {adresse}")
    qr_erzeugen(adresse, qr_datei)
    qr_pruefen(qr_datei, adresse)
    print(f"  gespeichert: {qr_datei.relative_to(WURZEL)}")
    return qr_datei, adresse


def kartensaetze():
    return sorted(d.name for d in (WURZEL / "docs").iterdir() if (d / "karten.js").exists())


def sammelblatt(ziel):
    klassen, titel = einstellungen_lesen()
    doc = Document()
    for abschnitt in doc.sections:
        abschnitt.left_margin = abschnitt.right_margin = Cm(2)
        abschnitt.top_margin = abschnitt.bottom_margin = Cm(1.8)
    kopf = doc.add_paragraph()
    r = kopf.add_run("QR-Codes der Lernkarten je Klasse")
    r.bold = True
    r.font.size = Pt(16)
    doc.add_paragraph(
        "Jeden Kasten an das Ende der passenden Ich-kann-Liste kopieren (Kasten anklicken → "
        "Tabelle markieren → Kopieren). Wer den Code scannt, landet automatisch in den "
        "Highscore-Listen seiner Klasse."
    )
    for ordner in kartensaetze():
        for kennung, name in klassen.items():
            qr_datei, adresse = qr_fuer(ordner, kennung)
            u = doc.add_paragraph()
            u.paragraph_format.space_before = Pt(14)
            u.paragraph_format.keep_with_next = True  # Überschrift nicht vom Kasten trennen
            r = u.add_run(f"{titel.get(ordner, ordner)} · Klasse {name}")
            r.bold = True
            r.font.size = Pt(12)
            kasten_anhaengen(doc, qr_datei, adresse, name)
    doc.save(ziel)
    print(f"Sammelblatt: {ziel}")


def main():
    a = argparse.ArgumentParser(description="QR-Codes für die Lernkarten erzeugen")
    a.add_argument("ordner", nargs="?", help="Ordner des Kartensatzes, z. B. kb1-diabetes")
    a.add_argument("liste", nargs="?", help="Ich-kann-Liste (.docx), in die der QR-Kasten soll")
    a.add_argument("--klasse", help="Kennung der Klasse aus einstellungen.js, z. B. 3bfp2-1")
    a.add_argument("--sammelblatt", metavar="ZIEL.docx", help="alle Kartensätze × alle Klassen in ein Dokument")
    arg = a.parse_args()

    if arg.sammelblatt:
        sammelblatt(arg.sammelblatt)
        return
    if not arg.ordner:
        sys.exit(__doc__)
    ordner = arg.ordner.strip("/")
    if not (WURZEL / "docs" / ordner / "karten.js").exists():
        sys.exit(f"docs/{ordner}/karten.js gibt es nicht.")
    klassenname = None
    if arg.klasse:
        klassen, _ = einstellungen_lesen()
        if arg.klasse not in klassen:
            sys.exit(f"Klasse „{arg.klasse}“ steht nicht in docs/app/einstellungen.js ({', '.join(klassen)}).")
        klassenname = klassen[arg.klasse]
    qr_datei, adresse = qr_fuer(ordner, arg.klasse)
    if arg.liste:
        print(f"  Ich-kann-Liste mit QR-Code: {in_liste_einsetzen(arg.liste, qr_datei, adresse, klassenname)}")


if __name__ == "__main__":
    main()
